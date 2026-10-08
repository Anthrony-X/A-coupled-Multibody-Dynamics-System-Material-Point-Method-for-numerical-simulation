from __future__ import annotations

import math
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "MBD-MPM模型建立方法小节.docx"
FIG_DIR = ROOT / "assets" / "generated_mbd_mpm_section"
XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")


def font_path(*names: str) -> str:
    for name in names:
        path = Path(r"C:\Windows\Fonts") / name
        if path.exists():
            return str(path)
    return str(Path(r"C:\Windows\Fonts\simsun.ttc"))


FONT_CN = font_path("msyh.ttc", "simhei.ttf", "simsun.ttc")
FONT_CN_BOLD = font_path("msyhbd.ttc", "simhei.ttf", "msyh.ttc")
FONT_EN = font_path("times.ttf", "calibri.ttf", "arial.ttf")


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_CN_BOLD if bold else FONT_CN, size=size)


def draw_centered(draw: ImageDraw.ImageDraw, box, text: str, font, fill=(30, 30, 30), spacing=8):
    x0, y0, x1, y1 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = y0 + ((y1 - y0) - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        x = x0 + ((x1 - x0) - w) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += h + spacing


def arrow(draw, start, end, fill, width=6, head=22):
    x0, y0 = start
    x1, y1 = end
    draw.line((x0, y0, x1, y1), fill=fill, width=width)
    ang = math.atan2(y1 - y0, x1 - x0)
    p1 = (x1 - head * math.cos(ang - math.pi / 6), y1 - head * math.sin(ang - math.pi / 6))
    p2 = (x1 - head * math.cos(ang + math.pi / 6), y1 - head * math.sin(ang + math.pi / 6))
    draw.polygon([end, p1, p2], fill=fill)


def round_rect(draw, box, radius, fill, outline, width=4):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_figures() -> tuple[Path, Path, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    fig1 = FIG_DIR / "fig1_coupling_framework.png"
    fig2 = FIG_DIR / "fig2_track_discretization.png"
    fig3 = FIG_DIR / "fig3_di_grid_contact.png"

    # Figure 1: coupled framework.
    scale = 2
    img = Image.new("RGB", (2200 * scale, 940 * scale), "white")
    draw = ImageDraw.Draw(img)
    title = pil_font(42 * scale, True)
    hfont = pil_font(33 * scale, True)
    bfont = pil_font(27 * scale)
    small = pil_font(23 * scale)
    accent = (50, 91, 132)
    green = (51, 118, 92)
    ochre = (155, 112, 40)
    grid = (235, 241, 247)
    for x in range(80 * scale, 2120 * scale, 80 * scale):
        draw.line((x, 140 * scale, x, 860 * scale), fill=grid, width=1)
    for y in range(160 * scale, 860 * scale, 80 * scale):
        draw.line((80 * scale, y, 2120 * scale, y), fill=grid, width=1)
    draw.text((90 * scale, 55 * scale), "MBD-MPM履带车辆-土体耦合建模框架", font=title, fill=(22, 54, 82))

    boxes = [
        ((120, 220, 610, 470), "坦克多体系统\n车体自由基座、悬架、车轮、\n主动轮/诱导轮、左右履带相位", (232, 242, 252), accent),
        ((855, 220, 1345, 470), "履带接触面离散\n履带鞋外表面 -> 接触片\n接触片 -> DI离散点", (239, 247, 242), green),
        ((1585, 220, 2075, 470), "MPM土体域\n物质点-背景网格映射\nDrucker-Prager应力更新", (252, 247, 235), ochre),
        ((855, 615, 1345, 805), "共享网格节点接触\n动量交换 + 摩擦/Janosi剪切\n输出左右履带接触反力", (246, 247, 249), (95, 95, 105)),
    ]
    for box, text, fill, outline in boxes:
        b = tuple(v * scale for v in box)
        round_rect(draw, b, 24 * scale, fill, outline, 5 * scale)
        draw_centered(draw, b, text, hfont if "\n" not in text else bfont, fill=(26, 42, 54), spacing=10 * scale)

    arrow(draw, (610 * scale, 345 * scale), (855 * scale, 345 * scale), accent, 8 * scale, 28 * scale)
    arrow(draw, (1345 * scale, 345 * scale), (1585 * scale, 345 * scale), green, 8 * scale, 28 * scale)
    arrow(draw, (1830 * scale, 470 * scale), (1345 * scale, 690 * scale), ochre, 7 * scale, 26 * scale)
    arrow(draw, (855 * scale, 690 * scale), (365 * scale, 470 * scale), (90, 90, 105), 7 * scale, 26 * scale)
    arrow(draw, (1100 * scale, 615 * scale), (1100 * scale, 470 * scale), (85, 120, 95), 7 * scale, 26 * scale)
    draw.text((676 * scale, 306 * scale), "姿态与履带鞋位姿", font=small, fill=accent)
    draw.text((1390 * scale, 306 * scale), "DI节点映射", font=small, fill=green)
    draw.text((1425 * scale, 585 * scale), "土体节点质量/速度", font=small, fill=ochre)
    draw.text((420 * scale, 603 * scale), "接触反力反馈", font=small, fill=(90, 90, 105))
    draw.text((96 * scale, 865 * scale), "Δt_MPM 为土体子步长；Δt_MBD = n_sub Δt_MPM 为多体宏步长。", font=small, fill=(80, 80, 88))
    img = img.resize((2200, 940), Image.Resampling.LANCZOS)
    img.save(fig1, quality=95)

    # Figure 2: side-view track discretization.
    img = Image.new("RGB", (2400 * scale, 1180 * scale), "white")
    draw = ImageDraw.Draw(img)
    title = pil_font(42 * scale, True)
    lab = pil_font(27 * scale)
    lab_b = pil_font(30 * scale, True)
    draw.text((90 * scale, 55 * scale), "履带包络路径与履带鞋离散", font=title, fill=(22, 54, 82))
    ground_y = 965 * scale
    draw.rectangle((80 * scale, ground_y, 2320 * scale, 1060 * scale), fill=(234, 225, 205))
    for x in range(100 * scale, 2320 * scale, 90 * scale):
        draw.line((x, ground_y, x + 55 * scale, 1060 * scale), fill=(210, 198, 176), width=2 * scale)
    # Hull.
    hull = [(520 * scale, 295 * scale), (1770 * scale, 265 * scale), (1960 * scale, 500 * scale), (360 * scale, 560 * scale)]
    draw.polygon(hull, fill=(215, 226, 233), outline=(52, 85, 108))
    draw.line(hull + [hull[0]], fill=(52, 85, 108), width=5 * scale)
    draw.polygon([(855 * scale, 260 * scale), (1450 * scale, 240 * scale), (1605 * scale, 315 * scale), (755 * scale, 350 * scale)], fill=(205, 218, 228), outline=(52, 85, 108))
    draw.line((1300 * scale, 270 * scale, 2040 * scale, 190 * scale), fill=(52, 85, 108), width=14 * scale)
    # Wheels.
    wheel_centers = [(520, 790), (765, 805), (1010, 810), (1255, 805), (1500, 790), (1745, 760)]
    for cx, cy in wheel_centers:
        draw.ellipse(((cx - 85) * scale, (cy - 85) * scale, (cx + 85) * scale, (cy + 85) * scale), fill=(245, 248, 250), outline=(55, 87, 112), width=6 * scale)
        draw.ellipse(((cx - 30) * scale, (cy - 30) * scale, (cx + 30) * scale, (cy + 30) * scale), fill=(92, 120, 142))
    sprocket = (330, 735, 105)
    idler = (1970, 705, 92)
    for cx, cy, r in [sprocket, idler]:
        draw.ellipse(((cx - r) * scale, (cy - r) * scale, (cx + r) * scale, (cy + r) * scale), fill=(247, 247, 243), outline=(88, 87, 78), width=6 * scale)
        for a in range(0, 360, 30):
            x0 = cx * scale
            y0 = cy * scale
            x1 = (cx + (r - 12) * math.cos(math.radians(a))) * scale
            y1 = (cy + (r - 12) * math.sin(math.radians(a))) * scale
            draw.line((x0, y0, x1, y1), fill=(111, 110, 99), width=3 * scale)
    for cx, cy in [(745, 610), (1150, 590), (1555, 610)]:
        draw.ellipse(((cx - 50) * scale, (cy - 50) * scale, (cx + 50) * scale, (cy + 50) * scale), fill=(247, 247, 243), outline=(88, 87, 78), width=5 * scale)

    # Track envelope approximate polyline with shoes.
    top = [(365, 625), (720, 545), (1150, 520), (1580, 545), (1950, 610)]
    bottom = [(420, 860), (760, 885), (1110, 895), (1460, 875), (1845, 815)]
    pts = top + bottom[::-1] + [top[0]]
    pts_s = [(x * scale, y * scale) for x, y in pts]
    draw.line(pts_s, fill=(42, 67, 83), width=34 * scale, joint="curve")
    draw.line(pts_s, fill=(112, 141, 158), width=20 * scale, joint="curve")

    # Discrete shoes as small short marks.
    shoe_color = (43, 92, 126)
    # Segment helper.
    path_segments = list(zip(pts[:-1], pts[1:]))
    samples = []
    for (x0, y0), (x1, y1) in path_segments:
        seg_len = math.hypot(x1 - x0, y1 - y0)
        n = max(2, int(seg_len / 82))
        for i in range(n):
            a = (i + 0.5) / n
            samples.append((x0 + a * (x1 - x0), y0 + a * (y1 - y0), math.atan2(y1 - y0, x1 - x0)))
    for x, y, a in samples:
        dx = 34 * math.cos(a)
        dy = 34 * math.sin(a)
        nx = -14 * math.sin(a)
        ny = 14 * math.cos(a)
        poly = [
            ((x - dx - nx) * scale, (y - dy - ny) * scale),
            ((x + dx - nx) * scale, (y + dy - ny) * scale),
            ((x + dx + nx) * scale, (y + dy + ny) * scale),
            ((x - dx + nx) * scale, (y - dy + ny) * scale),
        ]
        draw.polygon(poly, fill=shoe_color, outline=(21, 56, 78))
    arrow(draw, (330 * scale, 735 * scale), (330 * scale, 610 * scale), (160, 74, 60), 6 * scale, 24 * scale)
    draw.text((175 * scale, 560 * scale), "主动轮\nω_s", font=lab, fill=(130, 58, 50))
    draw.text((1860 * scale, 560 * scale), "诱导轮", font=lab, fill=(74, 74, 67))
    draw.text((855 * scale, 925 * scale), "下支履带由履带鞋接触片组成；相邻履带鞋间距为实际节距 p*", font=lab_b, fill=(30, 62, 88))
    arrow(draw, (650 * scale, 925 * scale), (520 * scale, 870 * scale), (30, 62, 88), 5 * scale, 22 * scale)
    draw.text((1440 * scale, 980 * scale), "MPM土体表面", font=lab, fill=(126, 99, 62))
    draw.text((925 * scale, 665 * scale), "车体自由基座 q_b = [x, y, h, ψ, θ]", font=lab, fill=(52, 85, 108))
    img = img.resize((2400, 1180), Image.Resampling.LANCZOS)
    img.save(fig2, quality=95)

    # Figure 3: DI grid contact.
    img = Image.new("RGB", (2250 * scale, 1150 * scale), "white")
    draw = ImageDraw.Draw(img)
    title = pil_font(42 * scale, True)
    lab = pil_font(27 * scale)
    lab_b = pil_font(30 * scale, True)
    draw.text((90 * scale, 55 * scale), "履带DI离散点与MPM背景网格的共享节点接触", font=title, fill=(22, 54, 82))
    x0, y0 = 170 * scale, 180 * scale
    dx = 150 * scale
    for i in range(13):
        draw.line((x0 + i * dx, y0, x0 + i * dx, y0 + 750 * scale), fill=(213, 225, 235), width=3 * scale)
    for j in range(6):
        draw.line((x0, y0 + j * dx, x0 + 1800 * scale, y0 + j * dx), fill=(213, 225, 235), width=3 * scale)
    # Soil particles.
    soil = [(430, 785), (580, 735), (720, 820), (875, 760), (1040, 835), (1185, 780), (1350, 825), (1510, 760), (1680, 805)]
    for x, y in soil:
        draw.ellipse(((x - 26) * scale, (y - 26) * scale, (x + 26) * scale, (y + 26) * scale), fill=(173, 127, 70), outline=(116, 83, 45), width=3 * scale)
    # Track patch.
    cx, cy = 1010 * scale, 500 * scale
    angle = math.radians(-13)
    L, W = 950 * scale, 250 * scale
    ca, sa = math.cos(angle), math.sin(angle)
    def rot(pt):
        x, y = pt
        return (cx + ca * x - sa * y, cy + sa * x + ca * y)
    corners = [rot((-L / 2, -W / 2)), rot((L / 2, -W / 2)), rot((L / 2, W / 2)), rot((-L / 2, W / 2))]
    draw.polygon(corners, fill=(195, 217, 229), outline=(43, 92, 126))
    draw.line(corners + [corners[0]], fill=(43, 92, 126), width=6 * scale)
    for ix in range(7):
        for iy in range(3):
            px = -L / 2 + (ix + 0.5) * L / 7
            py = -W / 2 + (iy + 0.5) * W / 3
            x, y = rot((px, py))
            draw.ellipse((x - 12 * scale, y - 12 * scale, x + 12 * scale, y + 12 * scale), fill=(21, 93, 134), outline=(255, 255, 255), width=2 * scale)
    # Highlight shared node and arrows.
    node = (1040 * scale, 630 * scale)
    draw.ellipse((node[0] - 20 * scale, node[1] - 20 * scale, node[0] + 20 * scale, node[1] + 20 * scale), fill=(210, 64, 61), outline=(124, 38, 36), width=4 * scale)
    arrow(draw, (1020 * scale, 530 * scale), (960 * scale, 405 * scale), (51, 118, 92), 7 * scale, 26 * scale)
    draw.text((875 * scale, 360 * scale), "n_I", font=lab_b, fill=(51, 118, 92))
    arrow(draw, (755 * scale, 515 * scale), (610 * scale, 555 * scale), (43, 92, 126), 7 * scale, 26 * scale)
    draw.text((540 * scale, 505 * scale), "v_t", font=lab_b, fill=(43, 92, 126))
    arrow(draw, (1000 * scale, 770 * scale), (1135 * scale, 730 * scale), (150, 93, 45), 7 * scale, 26 * scale)
    draw.text((1155 * scale, 708 * scale), "v_s", font=lab_b, fill=(150, 93, 45))
    arrow(draw, (1110 * scale, 620 * scale), (1305 * scale, 555 * scale), (190, 55, 52), 8 * scale, 28 * scale)
    draw.text((1320 * scale, 510 * scale), "接触冲量 J_I", font=lab_b, fill=(190, 55, 52))
    draw.text((300 * scale, 965 * scale), "蓝点：履带接触片上的DI节点；棕点：MPM土体物质点；红点：同时携带履带与土体动量的共享背景网格节点。", font=lab, fill=(70, 70, 78))
    img = img.resize((2250, 1150), Image.Resampling.LANCZOS)
    img.save(fig3, quality=95)
    return fig1, fig2, fig3


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None, color=None):
    run.font.name = "Times New Roman"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'<w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/>'
        r'<w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'
    )
    tbl_pr.append(borders)


def set_table_widths(table, widths_cm: list[float]):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


class EquationBuilder:
    def __init__(self):
        if not XSL.exists():
            raise FileNotFoundError(f"MML2OMML.XSL not found: {XSL}")
        self.transform = etree.XSLT(etree.parse(str(XSL)))

    def omml(self, mathml: str):
        src = etree.fromstring(mathml.encode("utf-8"))
        result = self.transform(src)
        return deepcopy(result.getroot())


MATH_NS = 'xmlns="http://www.w3.org/1998/Math/MathML"'


def mml(inner: str) -> str:
    return f'<math {MATH_NS}><mrow>{inner}</mrow></math>'


def add_math_equation(doc: Document, eq: EquationBuilder, mathml: str, number: int):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)
    set_table_widths(table, [14.8, 1.2])
    left, right = table.rows[0].cells
    for cell in (left, right):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=35, bottom=35, start=35, end=35)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(eq.omml(mathml))
    p2 = right.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p2.add_run(f"({number})")
    set_run_font(r, 10.5)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_paragraph(doc: Document, text: str, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74) if first_line else None
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(5)
    for idx, part in enumerate(text.split("`")):
        r = p.add_run(part)
        set_run_font(r, 10.5, italic=(idx % 2 == 1))
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, 14 if level == 1 else 12.5, bold=True, color=RGBColor(27, 71, 108))
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, 9, color=RGBColor(80, 80, 80))
    return p


def add_picture(doc: Document, path: Path, width_cm: float, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_note_box(doc: Document, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [16.0])
    remove_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF5F9")
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, 10.5, bold=True, color=RGBColor(27, 71, 108))
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, 9.8)


def add_algorithm_table(doc: Document):
    add_heading(doc, "2.X.4 多速率耦合求解流程", 2)
    add_paragraph(
        doc,
        "由于MPM显式积分的稳定步长通常显著小于坦克多体动力学的特征步长，本文采用多速率弱耦合策略。"
        "在一个MBD宏步内先根据上一宏步的平均接触反力推进坦克状态，再将宏步起点与终点的履带接触片进行时间插值，"
        "使MPM每个子步都获得连续的履带位置和速度。该处理避免了履带在MPM网格上出现阶跃式穿越，也便于在工程尺度土槽中保持较高计算效率。",
    )
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_widths(table, [1.4, 4.3, 10.3])
    headers = ["序号", "步骤", "主要操作"]
    for c, h in zip(table.rows[0].cells, headers):
        set_cell_shading(c, "DCEAF3")
        set_cell_margins(c, top=120, bottom=120, start=100, end=100)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 9.5, bold=True, color=RGBColor(27, 71, 108))
    rows = [
        ("1", "土体初始化", "生成MPM物质点并完成重力地应力平衡；保存土体位置、速度、仿射速度场和应力。"),
        ("2", "MBD宏步预测", "在 Δt_MBD = n_sub Δt_MPM 内，用上一宏步平均接触反力、驱动转矩、重力和阻尼推进车体、悬架、主动轮角速度与履带相位。"),
        ("3", "履带片插值", "对宏步起点和终点的接触片中心、法向和局部轴进行插值，计算每个MPM子步的DI节点位置和速度。"),
        ("4", "MPM P2G", "将土体物质点质量、动量和应力项映射到背景网格，施加重力和边界条件。"),
        ("5", "履带映射", "将DI节点质量、动量、法向和代表面积映射到相同背景网格节点，并记录左右履带贡献。"),
        ("6", "共享节点接触", "在同时含有土体和履带动量的节点上计算法向与切向冲量，更新土体网格速度并累计履带反力。"),
        ("7", "MPM G2P", "将修正后的网格速度回传至土体物质点，完成应变率、应力返回映射和物质点位置更新。"),
        ("8", "反力回传", "对子步履带反力按左右履带取平均，作为下一MBD宏步的外部接触载荷。"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for c, text in zip(cells, row):
            set_cell_margins(c, top=110, bottom=110, start=100, end=100)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in cells[:2] else WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            set_run_font(r, 9.2)


def add_initialization_and_checks(doc: Document, eq: EquationBuilder):
    add_heading(doc, "2.X.5 初始条件、边界条件与参数设置", 2)
    add_paragraph(
        doc,
        "为了降低履带车辆突然加载对土体显式积分造成的数值扰动，计算过程划分为地应力平衡、车辆沉降和驱动行驶三个阶段。"
        "地应力阶段仅求解MPM土体，在土槽内生成规则扰动的物质点云，并对土体施加重力及K0初始应力场。"
        "设土体上表面高度为zs，物质点竖向坐标为zp，则竖向应力和水平应力初始化为：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>σ</mi><mi>v</mi></msub><mo>=</mo><mo>−</mo><mi>ρ</mi><mi>g</mi><mo>(</mo><msub><mi>z</mi><mi>s</mi></msub><mo>−</mo><msub><mi>z</mi><mi>p</mi></msub><mo>)</mo>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>σ</mi><mi>h</mi></msub><mo>=</mo><msub><mi>K</mi><mn>0</mn></msub><msub><mi>σ</mi><mi>v</mi></msub>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>K</mi><mn>0</mn></msub><mo>=</mo><mn>1</mn><mo>−</mo><mi>sin</mi><mi>φ</mi>'
        ),
        24,
    )
    add_paragraph(
        doc,
        "沉降阶段将已平衡的土体状态作为初始场，坦克置于土体表面并逐步施加车体自重。"
        "该阶段锁定主动轮角速度、负重轮转动和履带相位，以表示制动车辆的静置沉降过程；同时可采用土体速度阻尼和车体垂向/俯仰阻尼吸收初始瞬态动能。"
        "自重加载系数采用线性斜坡：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>λ</mi><mi>g</mi></msub><mo>(</mo><mi>n</mi><mo>)</mo><mo>=</mo>'
            '<mi>min</mi><mo>(</mo><mn>1</mn><mo>,</mo><mfrac><mi>n</mi><msub><mi>N</mi><mi>ramp</mi></msub></mfrac><mo>)</mo>'
        ),
        25,
    )
    add_paragraph(
        doc,
        "行驶阶段从沉降后的车辆和土体状态继续计算，释放履带相位并施加左右主动轮转矩。"
        "MPM土槽底部节点施加固定边界，侧向和端部边界限制法向外流速度；顶部保持自由表面。"
        "背景网格采用包围土槽并带有外延距离的矩形计算域，以避免空网格过多造成计算浪费。",
    )

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_widths(table, [4.0, 4.0, 3.0, 5.0])
    headers = ["类别", "参数", "默认值", "作用"]
    for c, h in zip(table.rows[0].cells, headers):
        set_cell_shading(c, "DCEAF3")
        set_cell_margins(c, top=120, bottom=120, start=100, end=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 9.2, bold=True, color=RGBColor(27, 71, 108))
    rows = [
        ("土槽", "尺寸", "8.0 m x 5.0 m x 0.5 m", "对应前进、横向和竖向范围"),
        ("MPM", "物质点数", "96 x 60 x 8", "规则采样并叠加小扰动"),
        ("MPM", "背景网格", "n_grid = 96", "矩形域中按最大边长确定网格间距"),
        ("时间积分", "MPM子步长", "1.0e-5 s", "显式MPM积分步长"),
        ("时间积分", "MBD宏步长", "50个MPM子步", "与MPM步长取整数倍同步"),
        ("土体", "ρ, E, ν", "1700 kg/m3, 1.5 MPa, 0.30", "密度和线弹性参数"),
        ("土体", "φ, ψ, c", "18°, 2°, 4 kPa", "Drucker-Prager强度与剪胀参数"),
        ("接触", "μ, K", "0.50, 0.025 m", "摩擦系数与Janosi剪切位移模量"),
        ("履带DI", "节点间距/上限", "hg / 64", "默认以网格间距布置，每片最多64点"),
        ("车辆", "质量/驱动转矩", "42000 kg / 12000 N·m", "整车质量和单侧默认驱动输入"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, row)):
            set_cell_margins(cell, top=90, bottom=90, start=90, end=90)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2) else WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            set_run_font(r, 8.8)

    add_heading(doc, "2.X.6 数值稳定性与结果检查", 2)
    add_paragraph(
        doc,
        "显式MPM时间步长由网格尺寸、土体弹性波速、最大物质点速度和接触运动速度共同约束。"
        "在参数分析中应保证MPM子步长满足CFL型限制，并使MBD宏步长为MPM子步长的整数倍：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<mi>Δt</mi><mo>≤</mo><msub><mi>C</mi><mi>CFL</mi></msub>'
            '<mi>min</mi><mo>[</mo>'
            '<mfrac><msub><mi>h</mi><mi>g</mi></msub><msub><mi>c</mi><mi>s</mi></msub></mfrac>'
            '<mo>,</mo><mfrac><msub><mi>h</mi><mi>g</mi></msub><mrow><mo>|</mo><msub><mi mathvariant="bold">v</mi><mi>max</mi></msub><mo>|</mo><mo>+</mo><mi>ε</mi></mrow></mfrac>'
            '<mo>]</mo><mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>c</mi><mi>s</mi></msub><mo>=</mo>'
            '<msqrt><mfrac><mrow><mi>K</mi><mo>+</mo><mn>4</mn><mi>G</mi><mo>/</mo><mn>3</mn></mrow><mi>ρ</mi></mfrac></msqrt>'
        ),
        26,
    )
    add_paragraph(
        doc,
        "其中CCFL为小于1的安全系数，K和G分别为土体体积模量和剪切模量。"
        "接触处理的守恒性通过节点冲量的成对施加保证：土体网格速度增量所对应的冲量与反馈给履带多体系统的冲量大小相等、方向相反。"
        "实际计算中建议同步输出左右履带接触节点数、左右履带反力、车辆沉陷、俯仰角、前进速度、最大土体速度和土体自由表面高度，用于监测接触连续性和数值稳定性。",
    )
    add_paragraph(
        doc,
        "模型验证可从三个层次进行。首先，在无驱动沉降阶段比较车辆静态沉陷量、左右履带反力和总车重是否平衡；"
        "其次，在低速直线行驶阶段比较履带滑转率、牵引力和车辙深度随驱动转矩的变化趋势；"
        "最后，通过网格和DI节点间距敏感性分析检查解的收敛性。对任一输出量R，可采用相邻两级离散结果的相对差作为收敛指标：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>e</mi><mi>R</mi></msub><mo>=</mo>'
            '<mfrac><mrow><mo>|</mo><mi>R</mi><mo>(</mo><msub><mi>h</mi><mi>g</mi></msub><mo>)</mo><mo>−</mo><mi>R</mi><mo>(</mo><msub><mi>h</mi><mi>g</mi></msub><mo>/</mo><mn>2</mn><mo>)</mo><mo>|</mo></mrow>'
            '<mrow><mo>|</mo><mi>R</mi><mo>(</mo><msub><mi>h</mi><mi>g</mi></msub><mo>/</mo><mn>2</mn><mo>)</mo><mo>|</mo><mo>+</mo><mi>ε</mi></mrow></mfrac>'
        ),
        27,
    )
    add_paragraph(
        doc,
        "当车辆沉陷、平均牵引力和最大土体速度等关键响应随网格加密变化较小，且接触节点数随时间变化连续时，可认为履带离散和MPM网格分辨率满足当前工况的精度要求。",
    )


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.1)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        style.font.color.rgb = RGBColor(27, 71, 108)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("MBD-MPM模型建立方法")
    set_run_font(r, 8.5, color=RGBColor(110, 110, 110))


def build_docx():
    fig1, fig2, fig3 = make_figures()
    doc = Document()
    configure_document(doc)
    eq = EquationBuilder()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("2.X MBD-MPM模型建立方法")
    set_run_font(r, 17, bold=True, color=RGBColor(18, 64, 101))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(12)
    r = sub.add_run("坦克多体系统、履带离散与履带-MPM共享网格接触")
    set_run_font(r, 10.5, color=RGBColor(80, 80, 80))

    add_paragraph(
        doc,
        "为了描述重型履带车辆在可变形土体上的行驶、沉陷与牵引响应，本文建立了坦克多体动力学"
        "（multi-body dynamics, MBD）与物质点法（material point method, MPM）的耦合模型。"
        "车辆子系统用于给出车体姿态、车轮运动和履带鞋接触面的瞬时几何；MPM子系统用于解析土体的大变形、塑性屈服和自由表面演化；"
        "两者通过履带离散点在背景网格节点上的动量交换实现接触耦合。整体建模流程如图1所示。",
    )
    add_picture(doc, fig1, 15.8, "图1  MBD-MPM履带车辆-土体耦合建模框架")

    add_heading(doc, "2.X.1 坦克多体动力学模型", 2)
    add_paragraph(
        doc,
        "坦克几何由ZTZ96车辆视觉网格识别并重构为树型多体系统。模型包含车体、炮塔/火炮等固定于车体的上装部件，"
        "以及左右两侧的主动轮、诱导轮、托带轮、负重轮和悬架站。根据零部件识别结果，模型共包含1个车体自由基座、12个负重轮、"
        "6个托带轮、2个主动轮、2个诱导轮、12个悬架站和左右2条履带回路；炮塔、火炮、舱盖及附属装甲在动力学中并入相应刚体质量或作为固定子体处理。"
        "车体基座保留纵向位移、横向位移、垂向沉陷、偏航和俯仰自由度，侧倾在当前土槽直行/小偏航工况中忽略。",
    )
    add_paragraph(
        doc,
        "多体模型内部采用车辆视觉模型坐标系，其中x轴为横向、y轴为垂向、z轴为前进方向；MPM土体域采用x轴前进、y轴横向、z轴垂向的工程坐标系。"
        "在生成接触片时对点和向量进行统一坐标变换：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi mathvariant="bold">x</mi><mi>MPM</mi></msub><mo>=</mo>'
            '<msup><mrow><mo>[</mo><msub><mi>z</mi><mi>MBD</mi></msub><mo>,</mo>'
            '<msub><mi>x</mi><mi>MBD</mi></msub><mo>,</mo>'
            '<msub><mi>y</mi><mi>MBD</mi></msub><mo>]</mo></mrow><mi>T</mi></msup>'
        ),
        1,
    )
    add_paragraph(
        doc,
        "车辆广义坐标可写为车体坐标、左右履带相位和悬架转角的组合：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<mi mathvariant="bold">q</mi><mo>=</mo>'
            '<msup><mrow><mo>[</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>,</mo><mi>h</mi><mo>,</mo>'
            '<mi>ψ</mi><mo>,</mo><mi>θ</mi><mo>,</mo>'
            '<msub><mi>φ</mi><mi>L</mi></msub><mo>,</mo><msub><mi>φ</mi><mi>R</mi></msub><mo>,</mo>'
            '<msub><mi>θ</mi><mrow><mi>s</mi><mo>,</mo><mn>1</mn></mrow></msub><mo>,</mo><mi>⋯</mi><mo>,</mo>'
            '<msub><mi>θ</mi><mrow><mi>s</mi><mo>,</mo><mn>12</mn></mrow></msub><mo>]</mo></mrow><mi>T</mi></msup>'
        ),
        2,
    )
    add_paragraph(
        doc,
        "其中x、y分别为车辆质心在MPM水平平面内的纵向和横向位置，h为车体垂向位移，ψ和θ分别为偏航角和俯仰角，"
        "φη为第η侧（η=L,R）履带沿回路的相位，θs,i为第i个悬架站等效转角。多体动力学的一般形式写为：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<mi mathvariant="bold">M</mi><mo>(</mo><mi mathvariant="bold">q</mi><mo>)</mo>'
            '<mover><mi mathvariant="bold">q</mi><mo>¨</mo></mover>'
            '<mo>+</mo><mi mathvariant="bold">C</mi><mo>(</mo><mi mathvariant="bold">q</mi><mo>,</mo>'
            '<mover><mi mathvariant="bold">q</mi><mo>˙</mo></mover><mo>)</mo>'
            '<mover><mi mathvariant="bold">q</mi><mo>˙</mo></mover>'
            '<mo>+</mo><mi mathvariant="bold">g</mi><mo>(</mo><mi mathvariant="bold">q</mi><mo>)</mo>'
            '<mo>=</mo><msub><mi mathvariant="bold">Q</mi><mi>drv</mi></msub>'
            '<mo>+</mo><msubsup><mi mathvariant="bold">J</mi><mi>c</mi><mi>T</mi></msubsup>'
            '<msub><mi mathvariant="bold">f</mi><mi>c</mi></msub>'
            '<mo>−</mo><msub><mi mathvariant="bold">Q</mi><mi>d</mi></msub>'
        ),
        3,
    )
    add_paragraph(
        doc,
        "式中M为车辆质量矩阵，C为速度相关项，g为重力项，Qdrv为主动轮驱动转矩等广义驱动，Jc为履带-土体接触点到广义坐标的雅可比矩阵，"
        "fc为由MPM接触求解返回的履带反力，Qd包括空气阻力、滚动阻力、主动轮阻尼和车体垂向/俯仰阻尼。"
        "在实现中，上式按低阶车辆动力学显式积分：MPM返回的左右履带接触反力先投影到车辆前进方向，再分别用于更新纵向速度、偏航角速度、垂向速度和俯仰角速度。",
    )
    add_paragraph(
        doc,
        "车体偏航和俯仰转动惯量由识别出的车体包络尺寸估计。考虑视觉网格质量分布与真实车辆内部载荷不同，引入0.82的惯量修正系数：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>I</mi><mi>ψ</mi></msub><mo>=</mo>'
            '<mfrac><mrow><mn>0.82</mn><mi>m</mi><mo>(</mo><msup><mi>L</mi><mn>2</mn></msup><mo>+</mo><msubsup><mi>B</mi><mi>e</mi><mn>2</mn></msubsup><mo>)</mo></mrow><mn>12</mn></mfrac>'
            '<mo>,</mo><mspace width="0.5em"/>'
            '<msub><mi>I</mi><mi>θ</mi></msub><mo>=</mo>'
            '<mfrac><mrow><mn>0.82</mn><mi>m</mi><mo>(</mo><msup><mi>L</mi><mn>2</mn></msup><mo>+</mo><msup><mi>H</mi><mn>2</mn></msup><mo>)</mo></mrow><mn>12</mn></mfrac>'
        ),
        4,
    )
    add_paragraph(
        doc,
        "其中m为车辆总质量，L、H分别为车体长度和高度，Be=max(B, |xL|+|xR|+bt)为考虑左右履带中心距和履带宽度后的等效车宽。"
        "悬架站被等效为沿负重轮垂向运动的弹簧-阻尼单元。对第i个悬架站，接触压缩量和悬架支反力可写为：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>δ</mi><mi>i</mi></msub><mo>=</mo><mi>max</mi><mo>(</mo><mn>0</mn><mo>,</mo>'
            '<msub><mi>a</mi><mi>i</mi></msub><msub><mi>θ</mi><mrow><mi>s</mi><mo>,</mo><mi>i</mi></mrow></msub><mo>)</mo>'
            '<mo>,</mo><mspace width="0.5em"/>'
            '<msub><mi>N</mi><mi>i</mi></msub><mo>=</mo><msub><mi>N</mi><mrow><mn>0</mn><mo>,</mo><mi>i</mi></mrow></msub>'
            '<mo>+</mo><msub><mi>k</mi><mi>i</mi></msub><msub><mi>δ</mi><mi>i</mi></msub>'
            '<mo>+</mo><msub><mi>c</mi><mi>i</mi></msub><mover><msub><mi>δ</mi><mi>i</mi></msub><mo>˙</mo></mover>'
        ),
        5,
    )
    add_paragraph(
        doc,
        "为避免单个悬架站在显式积分中产生非物理拉力，各站支反力在每个时间步内进一步投影到满足整车垂向力平衡和俯仰力矩平衡的非负集合：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<munder><mi>min</mi><mrow><msub><mi>N</mi><mi>i</mi></msub><mo>≥</mo><mn>0</mn></mrow></munder>'
            '<mo>∑</mo><msup><mrow><mo>(</mo><msub><mi>N</mi><mi>i</mi></msub><mo>−</mo><msub><mover><mi>N</mi><mo>~</mo></mover><mi>i</mi></msub><mo>)</mo></mrow><mn>2</mn></msup>'
            '<mspace width="0.6em"/><mi>s.t.</mi><mspace width="0.4em"/>'
            '<mo>∑</mo><msub><mi>N</mi><mi>i</mi></msub><mo>=</mo><mi>m</mi><mi>g</mi>'
            '<mo>,</mo><mspace width="0.4em"/><mo>∑</mo><msub><mi>N</mi><mi>i</mi></msub><msub><mi>z</mi><mi>i</mi></msub><mo>=</mo><msub><mi>M</mi><mi>θ</mi></msub>'
        ),
        6,
    )
    add_paragraph(
        doc,
        "主动轮子系统通过驱动转矩与履带反力耦合。对左右履带分别有：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>I</mi><mi>s</mi></msub><mover><msub><mi>ω</mi><mi>η</mi></msub><mo>˙</mo></mover>'
            '<mo>=</mo><mo>−</mo><msub><mi>T</mi><mi>η</mi></msub>'
            '<mo>+</mo><msub><mi>F</mi><mrow><mi>x</mi><mo>,</mo><mi>η</mi></mrow></msub><msub><mi>r</mi><mi>s</mi></msub>'
            '<mo>−</mo><msub><mi>c</mi><mi>s</mi></msub><msub><mi>ω</mi><mi>η</mi></msub>'
            '<mo>,</mo><mspace width="0.5em"/>'
            '<msub><mi>v</mi><mrow><mi>t</mi><mo>,</mo><mi>η</mi></mrow></msub><mo>=</mo><mo>−</mo><msub><mi>r</mi><mi>s</mi></msub><msub><mi>ω</mi><mi>η</mi></msub>'
        ),
        7,
    )

    add_heading(doc, "2.X.2 履带离散与运动学描述", 2)
    add_paragraph(
        doc,
        "履带并不直接采用原始OBJ网格的细碎面片参与接触，而是由主动轮、诱导轮、托带轮和负重轮外包络构成闭合履带路径Γ。"
        "路径由圆弧段和相邻轮系之间的外公切线段组成，按弧长s进行参数化。该处理保留履带绕轮系运动的几何约束，同时避免将视觉网格中的非结构化三角面直接用于接触搜索。"
        "履带路径及履带鞋离散示意如图2所示。",
    )
    add_picture(doc, fig2, 16.0, "图2  履带包络路径、轮系与履带鞋离散示意")
    add_paragraph(
        doc,
        "对第η侧履带，若闭合路径总长为LΓ,η，名义节距为p，则履带鞋数量Nη和实际节距pη*取为：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>N</mi><mi>η</mi></msub><mo>=</mo><mi>round</mi><mo>(</mo>'
            '<mfrac><msub><mi>L</mi><mrow><mi>Γ</mi><mo>,</mo><mi>η</mi></mrow></msub><mi>p</mi></mfrac><mo>)</mo>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msubsup><mi>p</mi><mi>η</mi><mo>*</mo></msubsup><mo>=</mo>'
            '<mfrac><msub><mi>L</mi><mrow><mi>Γ</mi><mo>,</mo><mi>η</mi></mrow></msub><msub><mi>N</mi><mi>η</mi></msub></mfrac>'
        ),
        8,
    )
    add_paragraph(
        doc,
        "第i块履带鞋在路径上的弧长坐标由履带相位φη控制：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>s</mi><mrow><mi>i</mi><mo>,</mo><mi>η</mi></mrow></msub>'
            '<mo>=</mo><mi>i</mi><msubsup><mi>p</mi><mi>η</mi><mo>*</mo></msubsup><mo>−</mo><msub><mi>φ</mi><mi>η</mi></msub>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>φ</mi><mi>η</mi></msub><mo>(</mo><mi>t</mi><mo>+</mo><mi>Δt</mi><mo>)</mo>'
            '<mo>=</mo><msub><mi>φ</mi><mi>η</mi></msub><mo>(</mo><mi>t</mi><mo>)</mo>'
            '<mo>+</mo><msub><mi>v</mi><mrow><mi>t</mi><mo>,</mo><mi>η</mi></mrow></msub><mi>Δt</mi>'
        ),
        9,
    )
    add_paragraph(
        doc,
        "沿Γ采样可得到履带鞋中心的局部纵向-垂向坐标[yΓ(s), zΓ(s)]及切向tΓ(s)。"
        "再结合履带中心横向坐标xη、车体刚体变换Tb^w(q)以及局部轴定义，可得到第i块履带鞋的世界坐标位姿：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msubsup><mi mathvariant="bold">x</mi><mrow><mi>i</mi><mo>,</mo><mi>η</mi></mrow><mi>w</mi></msubsup>'
            '<mo>=</mo><msubsup><mi mathvariant="bold">T</mi><mi>b</mi><mi>w</mi></msubsup><mo>(</mo><mi mathvariant="bold">q</mi><mo>)</mo>'
            '<msup><mrow><mo>[</mo><msub><mi>x</mi><mi>η</mi></msub><mo>,</mo><msub><mi>y</mi><mi>Γ</mi></msub><mo>(</mo><msub><mi>s</mi><mrow><mi>i</mi><mo>,</mo><mi>η</mi></mrow></msub><mo>)</mo><mo>,</mo>'
            '<msub><mi>z</mi><mi>Γ</mi></msub><mo>(</mo><msub><mi>s</mi><mrow><mi>i</mi><mo>,</mo><mi>η</mi></mrow></msub><mo>)</mo><mo>]</mo></mrow><mi>T</mi></msup>'
        ),
        10,
    )
    add_paragraph(
        doc,
        "每个履带鞋在接触算法中被简化为一个刚性矩形接触片。接触片中心ci位于履带鞋外表面，外移距离为半履带厚度加履刺高度；"
        "局部轴li、bi和ni分别表示履带纵向、宽向和面法向，其中ni取指向土体的一侧。接触片半长ai=0.5 lshoe，半宽bi=0.5 btrack。"
        "为反映负重轮挤压导致的下支履带局部下挠，模型在下支段引入一维履带超单元：节点间以预张力和轴向刚度连接，轮-履带与地形/土体表面以罚函数支承，"
        "经若干松弛迭代后得到下支履带的变形剖面，并将该剖面回写到履带鞋姿态中。",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>T</mi><mi>e</mi></msub><mo>=</mo><mi>max</mi><mo>[</mo><mn>0</mn><mo>,</mo>'
            '<msub><mi>T</mi><mn>0</mn></msub><mo>+</mo><mi>E</mi><mi>A</mi>'
            '<mfrac><mrow><msub><mi>l</mi><mi>e</mi></msub><mo>−</mo><msub><mi>l</mi><mrow><mn>0</mn><mo>,</mo><mi>e</mi></mrow></msub></mrow>'
            '<msub><mi>l</mi><mrow><mn>0</mn><mo>,</mo><mi>e</mi></mrow></msub></mfrac><mo>]</mo>'
        ),
        11,
    )
    add_paragraph(
        doc,
        "式中T0为履带预张力，EA为履带等效轴向刚度，le和l0,e分别为当前单元长度与初始单元长度。该超单元只用于生成更合理的下支履带几何和轮-履带载荷分配；"
        "真实土体大变形仍由MPM计算。",
    )

    add_heading(doc, "2.X.3 履带离散点与MPM土体接触", 2)
    add_paragraph(
        doc,
        "在履带-MPM接触中，履带鞋接触片进一步离散为分布式相互作用（distributed interaction, DI）点。"
        "DI点不是土体物质点，也不参与土体本构积分；它们仅携带履带接触面的质量权重、速度、法向和代表面积，并映射到MPM背景网格。"
        "这种做法将复杂的刚体面接触转化为背景网格节点上的局部动量交换，适合处理履带与土体自由表面的频繁接触、脱离和再接触。图3给出了DI节点与MPM网格共享节点接触的示意。",
    )
    add_picture(doc, fig3, 15.8, "图3  履带DI节点与MPM背景网格共享节点接触示意")
    add_paragraph(
        doc,
        "对第i个接触片，沿履带纵向和宽向按目标间距ddi布置nli×nbi个DI点，并限制每个接触片的最大点数以控制计算量。第(j,k)个DI点的位置、质量和代表面积为：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi mathvariant="bold">x</mi><mrow><mi>i</mi><mi>j</mi><mi>k</mi></mrow></msub>'
            '<mo>=</mo><msub><mi mathvariant="bold">c</mi><mi>i</mi></msub>'
            '<mo>+</mo><msub><mi>ξ</mi><mi>j</mi></msub><msub><mi mathvariant="bold">l</mi><mi>i</mi></msub>'
            '<mo>+</mo><msub><mi>ζ</mi><mi>k</mi></msub><msub><mi mathvariant="bold">b</mi><mi>i</mi></msub>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>m</mi><mrow><mi>i</mi><mi>j</mi><mi>k</mi></mrow></msub><mo>=</mo>'
            '<mfrac><msub><mi>m</mi><mi>i</mi></msub><mrow><msub><mi>n</mi><mi>l</mi></msub><msub><mi>n</mi><mi>b</mi></msub></mrow></mfrac>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>A</mi><mrow><mi>i</mi><mi>j</mi><mi>k</mi></mrow></msub><mo>=</mo>'
            '<mfrac><mrow><mn>4</mn><msub><mi>a</mi><mi>i</mi></msub><msub><mi>b</mi><mi>i</mi></msub></mrow><mrow><msub><mi>n</mi><mi>l</mi></msub><msub><mi>n</mi><mi>b</mi></msub></mrow></mfrac>'
        ),
        12,
    )
    add_paragraph(
        doc,
        "DI点速度由当前接触片中心和上一子步接触片中心的差分确定。当MBD宏步长大于MPM子步长时，宏步起点和终点接触片拓扑一致时直接对接触片中心和局部轴进行线性插值；"
        "若由于履带鞋进入或离开接触区导致拓扑变化，则在每个MPM子步重新生成接触片，以保证接触面与实际履带姿态一致。",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi mathvariant="bold">v</mi><mi>i</mi></msub><mo>=</mo>'
            '<mfrac><mrow><msubsup><mi mathvariant="bold">c</mi><mi>i</mi><mrow><mi>n</mi><mo>+</mo><mn>1</mn></mrow></msubsup>'
            '<mo>−</mo><msubsup><mi mathvariant="bold">c</mi><mi>i</mi><mi>n</mi></msubsup></mrow><mi>Δt</mi></mfrac>'
        ),
        13,
    )
    add_paragraph(
        doc,
        "土体MPM采用二次B样条粒子-网格映射。对土体物质点p，其质量、动量和应力项映射到节点I的形式可概括为：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>m</mi><mi>I</mi></msub><mo>=</mo><munder><mo>∑</mo><mi>p</mi></munder><msub><mi>w</mi><mrow><mi>I</mi><mi>p</mi></mrow></msub><msub><mi>m</mi><mi>p</mi></msub>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi mathvariant="bold">p</mi><mi>I</mi></msub><mo>=</mo><munder><mo>∑</mo><mi>p</mi></munder><msub><mi>w</mi><mrow><mi>I</mi><mi>p</mi></mrow></msub>'
            '<mo>(</mo><msub><mi>m</mi><mi>p</mi></msub><msub><mi mathvariant="bold">v</mi><mi>p</mi></msub>'
            '<mo>+</mo><msub><mi mathvariant="bold">A</mi><mi>p</mi></msub><msub><mi mathvariant="bold">d</mi><mrow><mi>I</mi><mi>p</mi></mrow></msub><mo>)</mo>'
        ),
        14,
    )
    add_paragraph(
        doc,
        "其中wIp为MPM插值权重，Ap包含应力和仿射速度场贡献，dIp为粒子到网格节点的相对位置。土体本构采用Drucker-Prager屈服准则，"
        "由弹性预测和塑性返回映射更新物质点应力：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<mi>f</mi><mo>=</mo><msqrt><msub><mi>J</mi><mn>2</mn></msub></msqrt>'
            '<mo>−</mo><mi>α</mi><msub><mi>I</mi><mn>1</mn></msub><mo>−</mo><mi>k</mi><mo>≤</mo><mn>0</mn>'
        ),
        15,
    )
    add_paragraph(
        doc,
        "履带DI点映射到网格时采用单元内三线性权重，而非MPM物质点的二次核函数。原因是DI点代表刚体接触面上的几何采样点，"
        "应尽量保持接触作用局限于其所在背景单元。对网格节点I，履带侧质量、动量、法向和代表面积由下式累加：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msubsup><mi>m</mi><mi>I</mi><mi>t</mi></msubsup><mo>=</mo><munder><mo>∑</mo><mi>a</mi></munder><msub><mi>N</mi><mrow><mi>I</mi><mi>a</mi></mrow></msub><msub><mi>m</mi><mi>a</mi></msub>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msubsup><mi mathvariant="bold">p</mi><mi>I</mi><mi>t</mi></msubsup><mo>=</mo><munder><mo>∑</mo><mi>a</mi></munder><msub><mi>N</mi><mrow><mi>I</mi><mi>a</mi></mrow></msub><msub><mi>m</mi><mi>a</mi></msub><msub><mi mathvariant="bold">v</mi><mi>a</mi></msub>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi mathvariant="bold">n</mi><mi>I</mi></msub><mo>=</mo><mfrac><mrow><munder><mo>∑</mo><mi>a</mi></munder><msub><mi>N</mi><mrow><mi>I</mi><mi>a</mi></mrow></msub><msub><mi>m</mi><mi>a</mi></msub><msub><mi mathvariant="bold">n</mi><mi>a</mi></msub></mrow><mrow><mo>|</mo><munder><mo>∑</mo><mi>a</mi></munder><msub><mi>N</mi><mrow><mi>I</mi><mi>a</mi></mrow></msub><msub><mi>m</mi><mi>a</mi></msub><msub><mi mathvariant="bold">n</mi><mi>a</mi></msub><mo>|</mo></mrow></mfrac>'
        ),
        16,
    )
    add_paragraph(
        doc,
        "为避免履带面在尚未真正接近土体时因插值核扩散而提前接触，引入法向权重截断。设网格间距为hg，物质点等效间距dp=Vp1/3，默认截断权重取：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>w</mi><mi>cut</mi></msub><mo>=</mo><mn>1</mn><mo>−</mo>'
            '<mfrac><msub><mi>d</mi><mi>p</mi></msub><mrow><mn>2</mn><msub><mi>h</mi><mi>g</mi></msub></mrow></mfrac>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>d</mi><mi>p</mi></msub><mo>=</mo><msup><msub><mi>V</mi><mi>p</mi></msub><mfrac><mn>1</mn><mn>3</mn></mfrac></msup>'
        ),
        17,
    )
    add_paragraph(
        doc,
        "仅当某个网格节点同时携带土体质量和履带质量、法向权重超过阈值、且履带相对土体沿接触法向闭合时，才判定该节点发生接触：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<mi>I</mi><mo>∈</mo><mi>𝒞</mi><mo>⇔</mo>'
            '<msubsup><mi>m</mi><mi>I</mi><mi>s</mi></msubsup><mo>></mo><mn>0</mn><mo>,</mo>'
            '<msubsup><mi>m</mi><mi>I</mi><mi>t</mi></msubsup><mo>></mo><mn>0</mn><mo>,</mo>'
            '<msubsup><mi>w</mi><mi>I</mi><mi>max</mi></msubsup><mo>></mo><msub><mi>w</mi><mi>cut</mi></msub><mo>,</mo>'
            '<mo>(</mo><msubsup><mi mathvariant="bold">v</mi><mi>I</mi><mi>t</mi></msubsup><mo>−</mo><msubsup><mi mathvariant="bold">v</mi><mi>I</mi><mi>s</mi></msubsup><mo>)</mo><mo>·</mo><msub><mi mathvariant="bold">n</mi><mi>I</mi></msub><mo>></mo><mn>0</mn>'
        ),
        18,
    )
    add_paragraph(
        doc,
        "在接触节点I上，土体和履带的试算速度分别为vIs和vIt，相对速度为vrel=vIt−vIs。"
        "法向冲量由节点等效质量和闭合法向速度决定：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>m</mi><mi>eff</mi></msub><mo>=</mo>'
            '<mfrac><mrow><msubsup><mi>m</mi><mi>I</mi><mi>s</mi></msubsup><msubsup><mi>m</mi><mi>I</mi><mi>t</mi></msubsup></mrow><mrow><msubsup><mi>m</mi><mi>I</mi><mi>s</mi></msubsup><mo>+</mo><msubsup><mi>m</mi><mi>I</mi><mi>t</mi></msubsup></mrow></mfrac>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>J</mi><mi>n</mi></msub><mo>=</mo><msub><mi>m</mi><mi>eff</mi></msub>'
            '<mo>[</mo><mo>(</mo><msubsup><mi mathvariant="bold">v</mi><mi>I</mi><mi>t</mi></msubsup><mo>−</mo><msubsup><mi mathvariant="bold">v</mi><mi>I</mi><mi>s</mi></msubsup><mo>)</mo><mo>·</mo><msub><mi mathvariant="bold">n</mi><mi>I</mi></msub><mo>]</mo>'
        ),
        19,
    )
    add_paragraph(
        doc,
        "切向冲量先按粘着条件计算，再由界面强度截断。若采用库仑摩擦模型，切向冲量上限为μJn；"
        "若采用Janosi-Hanamoto界面剪切模型，则切向强度随累计剪切位移jI逐渐动员：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>J</mi><mi>t</mi></msub><mo>=</mo><mi>min</mi><mo>(</mo>'
            '<msub><mi>m</mi><mi>eff</mi></msub><mo>|</mo><msub><mi mathvariant="bold">v</mi><mi>τ</mi></msub><mo>|</mo>'
            '<mo>,</mo><msub><mi>J</mi><mrow><mi>t</mi><mo>,</mo><mi>lim</mi></mrow></msub><mo>)</mo>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msub><mi>J</mi><mrow><mi>t</mi><mo>,</mo><mi>lim</mi></mrow></msub><mo>=</mo>'
            '<mi>μ</mi><msub><mi>J</mi><mi>n</mi></msub>'
        ),
        20,
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msub><mi>J</mi><mrow><mi>t</mi><mo>,</mo><mi>lim</mi></mrow></msub><mo>=</mo>'
            '<mo>(</mo><msub><mi>c</mi><mi>a</mi></msub><mo>+</mo><msub><mi>σ</mi><mi>n</mi></msub><mi>tan</mi><msub><mi>φ</mi><mi>a</mi></msub><mo>)</mo>'
            '<msub><mi>A</mi><mi>I</mi></msub><mi>Δt</mi>'
            '<mo>[</mo><mn>1</mn><mo>−</mo><mi>exp</mi><mo>(</mo><mo>−</mo><mfrac><msub><mi>j</mi><mi>I</mi></msub><mi>K</mi></mfrac><mo>)</mo><mo>]</mo>'
        ),
        21,
    )
    add_paragraph(
        doc,
        "式中vτ为相对速度的切向分量，μ为履带-土体摩擦系数，ca和φa为界面黏聚力与界面摩擦角，"
        "σn=(Jn/Δt)/AI为节点法向接触应力，K为Janosi剪切位移模量。"
        "接触冲量只修正土体网格速度，履带作为外部多体系统不在MPM网格中积分，其反作用力通过牛顿第三定律累计后反馈给MBD：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<msubsup><mi mathvariant="bold">v</mi><mi>I</mi><mrow><mi>s</mi><mo>+</mo></mrow></msubsup><mo>=</mo>'
            '<msubsup><mi mathvariant="bold">v</mi><mi>I</mi><mrow><mi>s</mi><mo>−</mo></mrow></msubsup>'
            '<mo>+</mo><mfrac><mrow><msub><mi>J</mi><mi>n</mi></msub><msub><mi mathvariant="bold">n</mi><mi>I</mi></msub>'
            '<mo>+</mo><msub><mi>J</mi><mi>t</mi></msub><msub><mi mathvariant="bold">t</mi><mi>I</mi></msub></mrow><msubsup><mi>m</mi><mi>I</mi><mi>s</mi></msubsup></mfrac>'
            '<mo>,</mo><mspace width="0.6em"/>'
            '<msubsup><mi mathvariant="bold">F</mi><mi>I</mi><mi>t</mi></msubsup><mo>=</mo>'
            '<mo>−</mo><mfrac><mrow><msub><mi>J</mi><mi>n</mi></msub><msub><mi mathvariant="bold">n</mi><mi>I</mi></msub>'
            '<mo>+</mo><msub><mi>J</mi><mi>t</mi></msub><msub><mi mathvariant="bold">t</mi><mi>I</mi></msub></mrow><mi>Δt</mi></mfrac>'
        ),
        22,
    )
    add_paragraph(
        doc,
        "左右履带的接触节点按DI点映射质量的侧向贡献进行归属，并在一个MBD宏步内对子步接触反力取平均：",
    )
    add_math_equation(
        doc,
        eq,
        mml(
            '<mover><msub><mi mathvariant="bold">F</mi><mi>η</mi></msub><mo>¯</mo></mover><mo>=</mo>'
            '<mfrac><mn>1</mn><msub><mi>n</mi><mi>sub</mi></msub></mfrac>'
            '<msubsup><mo>∑</mo><mrow><mi>k</mi><mo>=</mo><mn>1</mn></mrow><msub><mi>n</mi><mi>sub</mi></msub></msubsup>'
            '<munder><mo>∑</mo><mrow><mi>I</mi><mo>∈</mo><msub><mi>𝒞</mi><mi>η</mi></msub></mrow></munder>'
            '<msubsup><mi mathvariant="bold">F</mi><mrow><mi>I</mi><mo>,</mo><mi>η</mi></mrow><mrow><mi>t</mi><mo>,</mo><mi>k</mi></mrow></msubsup>'
        ),
        23,
    )
    add_paragraph(
        doc,
        "平均反力再投影为车辆纵向牵引力、垂向支承力和偏航力矩。由于履带接触在MPM网格上完成，土体表面的大沉陷、推土隆起和履带局部脱空无需额外接触搜索即可自然反映在可接触节点集合𝒞中。",
    )

    add_algorithm_table(doc)
    add_initialization_and_checks(doc, eq)

    add_heading(doc, "符号与参数说明", 2)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_widths(table, [3.1, 7.1, 2.1, 3.7])
    headers = ["符号", "含义", "单位", "备注"]
    for c, h in zip(table.rows[0].cells, headers):
        set_cell_shading(c, "DCEAF3")
        set_cell_margins(c, top=120, bottom=120, start=100, end=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 9.5, bold=True, color=RGBColor(27, 71, 108))
    symbols = [
        ("q", "车辆广义坐标向量", "-", "含车体、履带相位和悬架变量"),
        ("LΓ, p, p*", "履带路径长度、名义节距、实际节距", "m", "由轮系包络路径确定"),
        ("ci, li, bi, ni", "履带接触片中心、纵向轴、宽向轴、法向", "m / -", "由履带鞋位姿生成"),
        ("ddi", "DI节点目标间距", "m", "默认可取MPM网格尺寸"),
        ("hg, dp", "MPM网格间距、物质点等效间距", "m", "用于法向权重截断"),
        ("Jn, Jt", "接触节点法向和切向冲量", "N·s", "由共享网格节点动量交换得到"),
        ("μ, ca, φa, K", "界面摩擦系数、黏聚力、摩擦角、Janosi剪切位移模量", "-", "用于切向接触强度"),
        ("Fη", "第η侧履带传递给MBD的平均接触反力", "N", "η=L,R"),
    ]
    for row in symbols:
        cells = table.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, row)):
            set_cell_margins(cell, top=95, bottom=95, start=100, end=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2) else WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            set_run_font(r, 9)

    doc.core_properties.title = "MBD-MPM模型建立方法"
    doc.core_properties.subject = "SCI论文方法小节"
    doc.core_properties.author = "Codex"
    doc.save(OUT)


if __name__ == "__main__":
    build_docx()
    print(OUT)
